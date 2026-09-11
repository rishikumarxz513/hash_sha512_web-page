"""
In-memory file metadata extraction.
Extracts internal creation/modification dates from images (EXIF), PDFs, DOCX,
XLSX, PPT/PPTX, DOC, video files, and EXE/DLL PE binaries.
"""

import io
import json
import os
import struct
import tempfile
from datetime import datetime
from typing import Optional

# ─── MIME Detection ───────────────────────────────────────────────────────────
try:
    import magic

    def detect_mime(content: bytes) -> str:
        return magic.from_buffer(content, mime=True)
except ImportError:
    import mimetypes

    def detect_mime(content: bytes, filename: str = "") -> str:
        guess, _ = mimetypes.guess_type(filename)
        return guess or "application/octet-stream"


# ─── Date Helpers ─────────────────────────────────────────────────────────────

def _parse_exif_date(val) -> Optional[datetime]:
    """Parse EXIF-style date string like '2024:01:15 10:30:00'."""
    if not val:
        return None
    try:
        s = str(val).strip().rstrip("\x00")
        return datetime.strptime(s, "%Y:%m:%d %H:%M:%S")
    except (ValueError, TypeError):
        return None


def _parse_pdf_date(val) -> Optional[datetime]:
    """Parse PDF date string like 'D:20240115103000+05'30'' or variants."""
    if not val:
        return None
    try:
        s = str(val).strip()
        if s.startswith("D:"):
            s = s[2:]
        # Take first 14 chars: YYYYMMDDHHmmSS
        s = s[:14].ljust(14, "0")
        return datetime.strptime(s, "%Y%m%d%H%M%S")
    except (ValueError, TypeError):
        return None


def _filetime_to_datetime(ft: int) -> Optional[datetime]:
    """Convert Windows FILETIME (100-ns ticks since 1601-01-01) to datetime."""
    if not ft or ft <= 0:
        return None
    try:
        EPOCH_DIFF = 116444736000000000  # 100-ns intervals between 1601 and 1970
        timestamp = (ft - EPOCH_DIFF) / 10_000_000
        return datetime.utcfromtimestamp(timestamp)
    except (OSError, OverflowError, ValueError):
        return None


# ─── Image Metadata (Pillow / EXIF) ──────────────────────────────────────────

def _extract_image(content: bytes) -> dict:
    from PIL import Image
    from PIL.ExifTags import TAGS

    img = Image.open(io.BytesIO(content))
    exif_raw = img._getexif() if hasattr(img, "_getexif") else None

    created = None
    modified = None
    extra = {
        "format": img.format,
        "size": f"{img.width}x{img.height}",
    }

    if exif_raw:
        exif = {TAGS.get(k, k): v for k, v in exif_raw.items()}
        created = _parse_exif_date(exif.get("DateTimeOriginal"))
        modified = _parse_exif_date(exif.get("DateTime"))
        if not created:
            created = _parse_exif_date(exif.get("DateTimeDigitized"))

        # Camera info
        if exif.get("Make"):
            extra["camera_make"] = str(exif["Make"]).strip().rstrip("\x00")
        if exif.get("Model"):
            extra["camera_model"] = str(exif["Model"]).strip().rstrip("\x00")
        if exif.get("LensModel"):
            extra["lens_model"] = str(exif["LensModel"]).strip().rstrip("\x00")

        # Shooting settings
        if exif.get("ISOSpeedRatings"):
            extra["iso"] = str(exif["ISOSpeedRatings"])
        if exif.get("FNumber"):
            try:
                fn = exif["FNumber"]
                extra["aperture"] = f"f/{float(fn):.1f}"
            except (TypeError, ValueError):
                extra["aperture"] = str(fn)
        if exif.get("FocalLength"):
            try:
                fl = exif["FocalLength"]
                extra["focal_length"] = f"{float(fl):.1f} mm"
            except (TypeError, ValueError):
                extra["focal_length"] = str(fl)
        if exif.get("ExposureTime"):
            try:
                et = exif["ExposureTime"]
                val = float(et)
                if val < 1:
                    extra["exposure_time"] = f"1/{int(1/val)}s"
                else:
                    extra["exposure_time"] = f"{val}s"
            except (TypeError, ValueError, ZeroDivisionError):
                extra["exposure_time"] = str(et)

        # Flash
        if "Flash" in exif:
            flash_val = exif["Flash"]
            if isinstance(flash_val, int):
                extra["flash"] = "Fired" if flash_val & 1 else "Not fired"
            else:
                extra["flash"] = str(flash_val)

        # GPS presence (privacy-safe)
        if exif.get("GPSInfo"):
            extra["gps_data"] = "Present"

        # Software
        if exif.get("Software"):
            extra["software"] = str(exif["Software"]).strip().rstrip("\x00")

        # Color space
        color_space_map = {1: "sRGB", 2: "Adobe RGB", 65535: "Uncalibrated"}
        if "ColorSpace" in exif:
            extra["color_space"] = color_space_map.get(exif["ColorSpace"], str(exif["ColorSpace"]))

        # Orientation
        orientation_map = {
            1: "Normal", 2: "Mirrored", 3: "Rotated 180°",
            4: "Mirrored + 180°", 5: "Mirrored + 270° CW",
            6: "Rotated 90° CW", 7: "Mirrored + 90° CW", 8: "Rotated 270° CW"
        }
        if "Orientation" in exif:
            extra["orientation"] = orientation_map.get(exif["Orientation"], str(exif["Orientation"]))

    return {
        "created_at_internal": created,
        "modified_at_internal": modified,
        "extra": extra,
    }


# ─── PDF Metadata (PyPDF2) ───────────────────────────────────────────────────

def _extract_pdf(content: bytes) -> dict:
    from PyPDF2 import PdfReader

    reader = PdfReader(io.BytesIO(content))
    info = reader.metadata

    created = None
    modified = None
    extra = {"pages": len(reader.pages)}

    if info:
        created = _parse_pdf_date(info.get("/CreationDate"))
        modified = _parse_pdf_date(info.get("/ModDate"))
        if info.get("/Author"):
            extra["author"] = str(info["/Author"])
        if info.get("/Producer"):
            extra["producer"] = str(info["/Producer"])
        if info.get("/Title"):
            extra["title"] = str(info["/Title"])
        if info.get("/Subject"):
            extra["subject"] = str(info["/Subject"])
        if info.get("/Creator"):
            extra["creator_app"] = str(info["/Creator"])
        if info.get("/Keywords"):
            extra["keywords"] = str(info["/Keywords"])

    return {
        "created_at_internal": created,
        "modified_at_internal": modified,
        "extra": extra,
    }


# ─── DOCX Metadata (python-docx) ─────────────────────────────────────────────

def _extract_docx(content: bytes) -> dict:
    from docx import Document

    doc = Document(io.BytesIO(content))
    props = doc.core_properties

    created = props.created if isinstance(props.created, datetime) else None
    modified = props.modified if isinstance(props.modified, datetime) else None

    extra = {}
    if props.author:
        extra["author"] = props.author
    if props.title:
        extra["title"] = props.title
    if props.subject:
        extra["subject"] = props.subject
    if props.keywords:
        extra["keywords"] = props.keywords
    if props.category:
        extra["category"] = props.category
    if props.last_modified_by:
        extra["last_modified_by"] = props.last_modified_by
    if props.revision is not None:
        extra["revision"] = str(props.revision)
    if props.description:
        extra["description"] = props.description

    return {
        "created_at_internal": created,
        "modified_at_internal": modified,
        "extra": extra,
    }


# ─── XLSX Metadata (openpyxl) ────────────────────────────────────────────────

def _extract_xlsx(content: bytes) -> dict:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    props = wb.properties

    created = props.created if isinstance(props.created, datetime) else None
    modified = props.modified if isinstance(props.modified, datetime) else None

    extra = {"sheets": wb.sheetnames}
    if props.creator:
        extra["author"] = props.creator
    if props.title:
        extra["title"] = props.title
    if props.subject:
        extra["subject"] = props.subject
    if props.description:
        extra["description"] = props.description
    if props.lastModifiedBy:
        extra["last_modified_by"] = props.lastModifiedBy
    wb.close()

    return {
        "created_at_internal": created,
        "modified_at_internal": modified,
        "extra": extra,
    }


# ─── PPT / PPTX Metadata (python-pptx + olefile) ────────────────────────────

def _extract_ppt(content: bytes) -> dict:
    """
    Modern .pptx → python-pptx core_properties.
    Legacy .ppt  → olefile OLE summary properties.
    Auto-detects format by trying python-pptx first (ZIP-based), falling back
    to olefile (OLE2 compound document).
    """
    # Try modern PPTX first
    try:
        from pptx import Presentation

        prs = Presentation(io.BytesIO(content))
        props = prs.core_properties

        created = props.created if isinstance(props.created, datetime) else None
        modified = props.modified if isinstance(props.modified, datetime) else None

        extra = {}
        if props.author:
            extra["author"] = props.author
        if props.title:
            extra["title"] = props.title
        extra["slides"] = len(prs.slides)
        if props.subject:
            extra["subject"] = props.subject
        if props.keywords:
            extra["keywords"] = props.keywords
        if props.category:
            extra["category"] = props.category
        if props.last_modified_by:
            extra["last_modified_by"] = props.last_modified_by
        if props.description:
            extra["description"] = props.description

        return {
            "created_at_internal": created,
            "modified_at_internal": modified,
            "extra": extra,
        }
    except Exception:
        pass

    # Fall back to legacy OLE (.ppt)
    return _extract_ole_properties(content, file_type="ppt")


def _extract_ole_properties(content: bytes, file_type: str = "ole") -> dict:
    """Extract metadata from OLE2 compound documents (.doc, .ppt, .xls)."""
    import olefile

    ole = olefile.OleFileIO(io.BytesIO(content))
    meta = ole.get_metadata()

    created = None
    modified = None
    extra = {"format": file_type.upper()}

    if meta.create_time and isinstance(meta.create_time, datetime):
        created = meta.create_time
    if meta.last_saved_time and isinstance(meta.last_saved_time, datetime):
        modified = meta.last_saved_time

    if meta.author:
        author = meta.author if isinstance(meta.author, str) else meta.author.decode("utf-8", errors="replace")
        extra["author"] = author
    if meta.title:
        title = meta.title if isinstance(meta.title, str) else meta.title.decode("utf-8", errors="replace")
        extra["title"] = title

    ole.close()

    return {
        "created_at_internal": created,
        "modified_at_internal": modified,
        "extra": extra,
    }


# ─── DOC Metadata (olefile) ─────────────────────────────────────────────────

def _extract_doc(content: bytes) -> dict:
    """Extract metadata from legacy .doc files using olefile."""
    return _extract_ole_properties(content, file_type="doc")


# ─── Video Metadata (hachoir) ────────────────────────────────────────────────

def _extract_video(content: bytes) -> dict:
    """
    Extract video metadata using hachoir.
    Uses a temporary file for reliable parsing (hachoir's BytesIO support
    is unreliable for video containers that require seeking).
    """
    from hachoir.parser import createParser
    from hachoir.metadata import extractMetadata
    from hachoir.core import config as hachoir_config
    hachoir_config.quiet = True

    # Write to a temp file — hachoir's createParser is far more reliable
    tmp_fd, tmp_path = tempfile.mkstemp(suffix=".vid")
    try:
        os.write(tmp_fd, content)
        os.close(tmp_fd)

        parser = createParser(tmp_path)
        if parser is None:
            return {
                "created_at_internal": None,
                "modified_at_internal": None,
                "extra": {"status": "Could not parse video container"},
            }

        with parser:
            meta = extractMetadata(parser)
            if meta is None:
                return {
                    "created_at_internal": None,
                    "modified_at_internal": None,
                    "extra": {"status": "No metadata found in video"},
                }

            extra = {}

            # Pull all fields from exportDictionary for rich metadata
            try:
                raw_dict = meta.exportDictionary()
                meta_section = raw_dict.get("Metadata", raw_dict)
                for key, value in meta_section.items():
                    clean_key = key.lower().replace(" ", "_").replace("-", "_")
                    if clean_key not in ("creation_date", "last_modification"):
                        extra[clean_key] = str(value)
            except Exception:
                pass

            # Duration
            duration = meta.getValues("duration")
            if duration:
                d = duration[0]
                total_secs = int(d.total_seconds()) if hasattr(d, "total_seconds") else 0
                hours, remainder = divmod(total_secs, 3600)
                mins, secs = divmod(remainder, 60)
                extra["duration"] = f"{hours:02d}:{mins:02d}:{secs:02d}" if hours else f"{mins:02d}:{secs:02d}"

            # Resolution
            width_vals = meta.getValues("width")
            height_vals = meta.getValues("height")
            if width_vals and height_vals:
                extra["resolution"] = f"{width_vals[0]}x{height_vals[0]}"

            # Bitrate
            bit_rate = meta.getValues("bit_rate")
            if bit_rate:
                br = bit_rate[0]
                if isinstance(br, (int, float)):
                    extra["bitrate"] = f"{int(br) // 1000} kbps"
                else:
                    extra["bitrate"] = str(br)

            # Video codec
            video_comp = meta.getValues("compression")
            if video_comp:
                extra["video_codec"] = str(video_comp[0])

            # Frame rate
            frame_rate = meta.getValues("frame_rate")
            if frame_rate:
                extra["frame_rate"] = str(frame_rate[0])

            # MIME type
            mime_vals = meta.getValues("mime_type")
            if mime_vals:
                extra["container_mime"] = str(mime_vals[0])

            # Creation date
            created = None
            modified = None
            creation_vals = meta.getValues("creation_date")
            if creation_vals and isinstance(creation_vals[0], datetime):
                created = creation_vals[0]
            mod_vals = meta.getValues("last_modification")
            if mod_vals and isinstance(mod_vals[0], datetime):
                modified = mod_vals[0]

            return {
                "created_at_internal": created,
                "modified_at_internal": modified,
                "extra": extra,
            }
    finally:
        try:
            os.unlink(tmp_path)
        except OSError:
            pass


# ─── EXE / DLL Metadata (pefile) ────────────────────────────────────────────

def _extract_exe(content: bytes) -> dict:
    """
    Parse a Windows PE (EXE/DLL) and extract version info + architecture.
    Uses pefile, all in-memory.
    """
    import pefile

    pe = pefile.PE(data=content, fast_load=True)

    extra = {}

    # Machine architecture
    machine = pe.FILE_HEADER.Machine
    arch_map = {
        0x14C: "x86 (32-bit)",
        0x8664: "x64 (64-bit)",
        0xAA64: "ARM64",
        0x1C0: "ARM",
    }
    extra["architecture"] = arch_map.get(machine, f"Unknown (0x{machine:X})")

    # Compilation timestamp
    created = None
    try:
        ts = pe.FILE_HEADER.TimeDateStamp
        if ts and ts > 0:
            created = datetime.utcfromtimestamp(ts)
    except (OSError, OverflowError, ValueError):
        pass

    # VS_VERSIONINFO resource
    pe.parse_data_directories(
        directories=[pefile.DIRECTORY_ENTRY["IMAGE_DIRECTORY_ENTRY_RESOURCE"]]
    )

    if hasattr(pe, "FileInfo"):
        for fi_list in pe.FileInfo:
            for fi in fi_list:
                if hasattr(fi, "StringTable"):
                    for st in fi.StringTable:
                        entries = st.entries
                        for key, val in entries.items():
                            k = key.decode("utf-8", errors="replace") if isinstance(key, bytes) else str(key)
                            v = val.decode("utf-8", errors="replace") if isinstance(val, bytes) else str(val)

                            if k == "FileVersion":
                                extra["file_version"] = v
                            elif k == "ProductName":
                                extra["product_name"] = v
                            elif k == "OriginalFilename":
                                extra["original_filename"] = v
                            elif k == "CompanyName":
                                extra["company_name"] = v

    pe.close()

    return {
        "created_at_internal": created,
        "modified_at_internal": None,
        "extra": extra,
    }


# ─── Originality Check ───────────────────────────────────────────────────────

def _check_originality(
    created: Optional[datetime],
    modified: Optional[datetime],
) -> Optional[bool]:
    """
    Returns:
        True  — file appears original (not modified after creation)
        False — file has been modified significantly after creation
        None  — cannot determine (missing dates)
    """
    if created is None or modified is None:
        return None
    # Modified BEFORE created → timestamp manipulation, not original
    if modified < created:
        return False
    diff_seconds = abs((modified - created).total_seconds())
    # If modification is > 60s after creation, flag as modified
    return diff_seconds <= 60


def compute_originality_score(created: Optional[datetime], modified: Optional[datetime], extra: dict):
    """
    Computes an originality percentage (0.0 to 100.0) based on timestamps, 
    metadata completeness, and producer authenticity.
    Returns (score, breakdown_dict) or (None, None).
    """
    # If we have no creation date, we can't do a full analysis reliable enough for a score.
    if created is None:
        return None, None

    score = 0.0
    breakdown = {
        "timestamp_gap": {"score": 0, "max": 90, "label": "Timestamp Gap"},
        "completeness": {"score": 0, "max": 5, "label": "Metadata Completeness"},
        "authenticity": {"score": 0, "max": 5, "label": "Producer Authenticity"}
    }

    # 1. Timestamp Gap (90 points)
    if modified is None:
        breakdown["timestamp_gap"]["score"] = 90  # Perfect if never modified
    elif modified < created:
        # Modified BEFORE created → suspicious timestamp manipulation
        breakdown["timestamp_gap"]["score"] = 10
    else:
        diff_seconds = abs((modified - created).total_seconds())
        if diff_seconds <= 60:
            breakdown["timestamp_gap"]["score"] = 90
        elif diff_seconds <= 3600:  # 1 hour
            breakdown["timestamp_gap"]["score"] = 70
        elif diff_seconds <= 86400:  # 1 day
            breakdown["timestamp_gap"]["score"] = 50
        elif diff_seconds <= 2592000:  # 30 days
            breakdown["timestamp_gap"]["score"] = 30
        else:
            breakdown["timestamp_gap"]["score"] = 10

    # 2. Completeness (5 points)
    completeness_fields = ["author", "title", "company_name", "producer", "format", "resolution", "slides", "pages"]
    found_fields = sum(1 for field in completeness_fields if extra.get(field))
    if found_fields >= 2:
        breakdown["completeness"]["score"] = 5
    elif found_fields == 1:
        breakdown["completeness"]["score"] = 2
    else:
        breakdown["completeness"]["score"] = 0 

    # 3. Authenticity (5 points)
    authentic_producers = ["microsoft", "adobe", "apple", "google", "canon", "nikon", "sony"]
    modifying_producers = ["libreoffice", "foxit", "gimp", "pdf24", "ilovepdf", "ffmpeg", "handbrake"]
    
    producer_str = str(extra.get("producer", "")).lower() + " " + str(extra.get("author", "")).lower() + " " + str(extra.get("company_name", "")).lower()
    
    if any(p in producer_str for p in modifying_producers):
        breakdown["authenticity"]["score"] = 0
    elif any(p in producer_str for p in authentic_producers):
        breakdown["authenticity"]["score"] = 5
    else:
        breakdown["authenticity"]["score"] = 3 # default assumption

    score = sum(b["score"] for b in breakdown.values())
    return score, breakdown


# ─── Main Dispatcher ─────────────────────────────────────────────────────────

MIME_MAP = {
    # Images
    "image/jpeg": _extract_image,
    "image/tiff": _extract_image,
    "image/png": _extract_image,
    # Documents
    "application/pdf": _extract_pdf,
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": _extract_docx,
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": _extract_xlsx,
    # PowerPoint
    "application/vnd.ms-powerpoint": _extract_ppt,
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": _extract_ppt,
    # Legacy Word
    "application/msword": _extract_doc,
    # Video
    "video/mp4": _extract_video,
    "video/x-matroska": _extract_video,
    "video/x-msvideo": _extract_video,
    "video/quicktime": _extract_video,
    # Windows executables
    "application/x-dosexec": _extract_exe,
    "application/x-msdownload": _extract_exe,
    "application/vnd.microsoft.portable-executable": _extract_exe,
}


def _datetime_to_str(dt: Optional[datetime]) -> Optional[str]:
    return dt.isoformat() if dt else None


def extract_metadata(content: bytes, filename: str) -> dict:
    """
    Extract metadata from file bytes.

    Returns dict with keys:
        mime_type: str
        created_at_internal: datetime | None
        modified_at_internal: datetime | None
        is_original: bool | None
        metadata_json: str (JSON)
        metadata_available: bool
    """
    try:
        mime = detect_mime(content) if "filename" not in detect_mime.__code__.co_varnames else detect_mime(content, filename)
    except Exception:
        mime = "application/octet-stream"

    extractor = MIME_MAP.get(mime)

    if not extractor:
        return {
            "mime_type": mime,
            "created_at_internal": None,
            "modified_at_internal": None,
            "is_original": None,
            "originality_score": None,
            "originality_breakdown": None,
            "metadata_json": json.dumps({"status": "Metadata Unavailable", "mime_type": mime}),
            "metadata_available": False,
        }

    try:
        result = extractor(content)
    except Exception as e:
        return {
            "mime_type": mime,
            "created_at_internal": None,
            "modified_at_internal": None,
            "is_original": None,
            "originality_score": None,
            "originality_breakdown": None,
            "metadata_json": json.dumps({
                "status": "Metadata Extraction Failed",
                "mime_type": mime,
                "error": str(e),
            }),
            "metadata_available": False,
        }

    created = result.get("created_at_internal")
    modified = result.get("modified_at_internal")
    is_original = _check_originality(created, modified)
    originality_score, originality_breakdown = compute_originality_score(created, modified, result.get("extra", {}))

    metadata_obj = {
        "mime_type": mime,
        "created_at_internal": _datetime_to_str(created),
        "modified_at_internal": _datetime_to_str(modified),
        "is_original": is_original,
        "originality_score": originality_score,
        "originality_breakdown": originality_breakdown,
        "extra": result.get("extra", {}),
    }

    return {
        "mime_type": mime,
        "created_at_internal": created,
        "modified_at_internal": modified,
        "is_original": is_original,
        "originality_score": originality_score,
        "originality_breakdown": json.dumps(originality_breakdown) if originality_breakdown else None,
        "metadata_json": json.dumps(metadata_obj, default=str),
        "metadata_available": True,
    }
